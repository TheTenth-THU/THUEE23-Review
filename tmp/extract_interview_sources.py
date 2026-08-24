from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


ROOT = Path(sys.argv[1]).resolve()
SOURCE_DIR = ROOT / "test"
OUTPUT_DIR = ROOT / "tmp" / "source_text"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def safe_name(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*]+', "_", name)


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def extract_docx(stream: str | io.BytesIO) -> str:
    doc = Document(stream)
    blocks: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table_index, table in enumerate(doc.tables, start=1):
        blocks.append(f"[TABLE {table_index}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def extract_pptx(stream: str | io.BytesIO) -> str:
    presentation = Presentation(stream)
    blocks: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        blocks.append(f"[SLIDE {slide_index}]")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    blocks.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(blocks)


def extract_pdf(stream: str | io.BytesIO) -> str:
    reader = PdfReader(stream)
    blocks: list[str] = []
    for page_index, page in enumerate(reader.pages, start=1):
        blocks.append(f"[PAGE {page_index}]")
        blocks.append(page.extract_text() or "")
    return "\n".join(blocks)


def extract_xlsx(stream: str | io.BytesIO) -> str:
    workbook = load_workbook(stream, data_only=True, read_only=True)
    blocks: list[str] = []
    for sheet in workbook.worksheets:
        blocks.append(f"[SHEET {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                blocks.append(" | ".join(values))
    return "\n".join(blocks)


def extract_bytes(name: str, data: bytes) -> str | None:
    suffix = Path(name).suffix.lower()
    stream = io.BytesIO(data)
    if suffix == ".docx":
        return extract_docx(stream)
    if suffix == ".pptx":
        return extract_pptx(stream)
    if suffix == ".pdf":
        return extract_pdf(stream)
    if suffix in {".xlsx", ".xlsm"}:
        return extract_xlsx(stream)
    if suffix in {".md", ".txt", ".csv", ".tsv", ".json", ".yaml", ".yml"}:
        return decode_text(data)
    return None


def extract_file(path: Path) -> list[tuple[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".zip":
        outputs: list[tuple[str, str]] = []
        with zipfile.ZipFile(path) as archive:
            for info in archive.infolist():
                if info.is_dir():
                    continue
                if Path(info.filename).suffix.lower() == ".pdf":
                    continue
                text = extract_bytes(info.filename, archive.read(info))
                if text is not None:
                    outputs.append((f"{path.name}::{info.filename}", text))
        return outputs
    text = extract_bytes(path.name, path.read_bytes())
    return [] if text is None else [(path.name, text)]


index_lines: list[str] = []
error_lines: list[str] = []
for source_path in sorted(SOURCE_DIR.iterdir(), key=lambda item: item.name):
    if not source_path.is_file():
        continue
    try:
        extracted = extract_file(source_path)
    except Exception as exc:
        error_lines.append(f"{source_path.name}\t{type(exc).__name__}: {exc}")
        continue
    for logical_name, text in extracted:
        output_name = safe_name(logical_name) + ".txt"
        output_path = OUTPUT_DIR / output_name
        output_path.write_text(text, encoding="utf-8")
        index_lines.append(f"{logical_name}\t{output_name}\t{len(text)}")

(OUTPUT_DIR / "INDEX.tsv").write_text("\n".join(index_lines) + "\n", encoding="utf-8")
(OUTPUT_DIR / "ERRORS.tsv").write_text("\n".join(error_lines) + "\n", encoding="utf-8")
print(f"Extracted {len(index_lines)} readable objects to {OUTPUT_DIR}; errors: {len(error_lines)}")
